"""Native, structure-aware parsers — one per format, each returning a CanonicalDoc.

Every parser self-registers via `register(...)` at import time. Heavy/optional deps
(python-docx, python-pptx, pdfplumber) are imported INSIDE the methods that use them,
so importing this module never fails when those libraries are absent.
"""

from __future__ import annotations

import re
from pathlib import Path

from app.parsing.base import register
from app.parsing.canonical import CanonicalBlock, CanonicalDoc, CanonicalSection

# A fenced-code delimiter: ``` or ~~~ optionally followed by an info string.
_FENCE = re.compile(r"^(```|~~~)(.*)$")
# An ATX heading: 1-6 leading '#' then a space then the title.
_ATX = re.compile(r"^(#{1,6})\s+(.*)$")
# A GFM table separator row: | --- | :--: | ---: | etc.
_TABLE_SEP = re.compile(r"^\s*\|?\s*:?-{1,}:?\s*(\|\s*:?-{1,}:?\s*)+\|?\s*$")


def _split_row(line: str) -> list[str]:
    """Split a GFM pipe-table row into trimmed cells, dropping edge pipes."""
    s = line.strip()
    if s.startswith("|"):
        s = s[1:]
    if s.endswith("|"):
        s = s[:-1]
    return [c.strip() for c in s.split("|")]


def _is_table_row(line: str) -> bool:
    return "|" in line and line.strip().startswith(("|", "")) and line.strip() != ""


def markdown_to_canonical(
    text: str,
    *,
    doc_title: str,
    source_format: str = "md",
    source_path: str | None = None,
) -> CanonicalDoc:
    """Parse markdown into a CanonicalDoc.

    - ATX headings (`#`..`######`) maintain a heading stack -> CanonicalSection.path.
    - Fenced ```/~~~ blocks -> CanonicalBlock(type="code").
    - GFM pipe tables (header row + `---` separator + body) -> CanonicalBlock(type="table").
    - Other lines group into paragraph blocks; a blank line separates paragraphs.
    """
    lines = text.splitlines()
    sections: list[CanonicalSection] = []
    heading_stack: list[tuple[int, str]] = []  # (level, title)
    current = CanonicalSection(path=[])
    sections.append(current)

    para: list[str] = []

    def flush_para() -> None:
        if para:
            joined = "\n".join(para).strip()
            if joined:
                current.blocks.append(CanonicalBlock(type="paragraph", text=joined))
            para.clear()

    def start_section() -> CanonicalSection:
        path = [title for _, title in heading_stack]
        sec = CanonicalSection(path=path)
        sections.append(sec)
        return sec

    i = 0
    n = len(lines)
    while i < n:
        line = lines[i]

        # Fenced code block.
        fence = _FENCE.match(line)
        if fence:
            flush_para()
            marker = fence.group(1)
            body: list[str] = []
            i += 1
            while i < n:
                close = _FENCE.match(lines[i])
                if close and close.group(1) == marker:
                    i += 1
                    break
                body.append(lines[i])
                i += 1
            current.blocks.append(CanonicalBlock(type="code", text="\n".join(body)))
            continue

        # ATX heading.
        heading = _ATX.match(line)
        if heading:
            flush_para()
            level = len(heading.group(1))
            title = heading.group(2).strip()
            while heading_stack and heading_stack[-1][0] >= level:
                heading_stack.pop()
            heading_stack.append((level, title))
            current = start_section()
            current.blocks.append(CanonicalBlock(type="heading", text=title, level=level))
            i += 1
            continue

        # GFM pipe table: a row with pipes followed by a separator row.
        if "|" in line and line.strip() and i + 1 < n and _TABLE_SEP.match(lines[i + 1]):
            flush_para()
            header = _split_row(line)
            rows: list[list[str]] = [header]
            i += 2  # skip header + separator
            while i < n and "|" in lines[i] and lines[i].strip():
                if _TABLE_SEP.match(lines[i]):
                    i += 1
                    continue
                rows.append(_split_row(lines[i]))
                i += 1
            tsv = "\n".join("\t".join(r) for r in rows)
            current.blocks.append(CanonicalBlock(type="table", text=tsv, rows=rows))
            continue

        # Blank line separates paragraphs.
        if not line.strip():
            flush_para()
            i += 1
            continue

        para.append(line)
        i += 1

    flush_para()

    # Drop a leading empty pre-heading section if it carries no content.
    sections = [s for s in sections if s.blocks or s.path]
    if not sections:
        sections = [CanonicalSection(path=[])]

    return CanonicalDoc(
        doc_title=doc_title,
        source_format=source_format,
        source_path=source_path,
        sections=sections,
    )


def _paragraphs_to_section(text: str, *, path: list[str] | None = None) -> CanonicalSection:
    """Split plain text into paragraph blocks on blank lines."""
    blocks: list[CanonicalBlock] = []
    for chunk in re.split(r"\n\s*\n", text):
        cleaned = chunk.strip()
        if cleaned:
            blocks.append(CanonicalBlock(type="paragraph", text=cleaned))
    return CanonicalSection(path=path or [], blocks=blocks)


class MarkdownParser:
    formats: tuple[str, ...] = ("md",)

    def parse(self, path: str | Path, *, doc_title: str | None = None) -> CanonicalDoc:
        p = Path(path)
        text = p.read_text(encoding="utf-8")
        return markdown_to_canonical(
            text,
            doc_title=doc_title or p.stem,
            source_format="md",
            source_path=str(p),
        )


class TxtParser:
    formats: tuple[str, ...] = ("txt",)

    def parse(self, path: str | Path, *, doc_title: str | None = None) -> CanonicalDoc:
        p = Path(path)
        text = p.read_text(encoding="utf-8")
        section = _paragraphs_to_section(text)
        return CanonicalDoc(
            doc_title=doc_title or p.stem,
            source_format="txt",
            source_path=str(p),
            sections=[section],
        )


class DocxParser:
    formats: tuple[str, ...] = ("docx",)

    def parse(self, path: str | Path, *, doc_title: str | None = None) -> CanonicalDoc:
        import docx  # python-docx

        p = Path(path)
        document = docx.Document(str(p))

        sections: list[CanonicalSection] = []
        heading_stack: list[tuple[int, str]] = []
        current = CanonicalSection(path=[])
        sections.append(current)

        def start_section() -> CanonicalSection:
            sec = CanonicalSection(path=[title for _, title in heading_stack])
            sections.append(sec)
            return sec

        for para in document.paragraphs:
            text = (para.text or "").strip()
            style = getattr(para.style, "name", "") or ""
            match = re.match(r"^Heading\s+(\d+)$", style)
            if match:
                level = int(match.group(1))
                title = text
                while heading_stack and heading_stack[-1][0] >= level:
                    heading_stack.pop()
                heading_stack.append((level, title))
                current = start_section()
                current.blocks.append(
                    CanonicalBlock(type="heading", text=title, level=level)
                )
                continue
            if text:
                current.blocks.append(CanonicalBlock(type="paragraph", text=text))

        for table in document.tables:
            rows: list[list[str]] = []
            for row in table.rows:
                rows.append([(cell.text or "").strip() for cell in row.cells])
            if rows:
                tsv = "\n".join("\t".join(r) for r in rows)
                current.blocks.append(
                    CanonicalBlock(type="table", text=tsv, rows=rows)
                )

        sections = [s for s in sections if s.blocks or s.path]
        if not sections:
            sections = [CanonicalSection(path=[])]

        return CanonicalDoc(
            doc_title=doc_title or p.stem,
            source_format="docx",
            source_path=str(p),
            sections=sections,
        )


class PptxParser:
    formats: tuple[str, ...] = ("pptx",)

    def parse(self, path: str | Path, *, doc_title: str | None = None) -> CanonicalDoc:
        from pptx import Presentation  # python-pptx

        p = Path(path)
        prs = Presentation(str(p))

        sections: list[CanonicalSection] = []
        for idx, slide in enumerate(prs.slides, start=1):
            title = ""
            title_shape = getattr(slide.shapes, "title", None)
            if title_shape is not None and getattr(title_shape, "has_text_frame", False):
                title = (title_shape.text or "").strip()

            section = CanonicalSection(path=[f"Slide {idx}: {title}"])
            seen_title = False
            for shape in slide.shapes:
                if getattr(shape, "has_table", False):
                    tbl = shape.table
                    rows: list[list[str]] = []
                    for row in tbl.rows:
                        rows.append([(cell.text or "").strip() for cell in row.cells])
                    if rows:
                        tsv = "\n".join("\t".join(r) for r in rows)
                        section.blocks.append(
                            CanonicalBlock(type="table", text=tsv, rows=rows)
                        )
                    continue
                if getattr(shape, "has_text_frame", False):
                    text = (shape.text_frame.text or "").strip()
                    if not text:
                        continue
                    if not seen_title and shape is title_shape:
                        section.blocks.append(
                            CanonicalBlock(type="heading", text=text, level=1)
                        )
                        seen_title = True
                        continue
                    for chunk in re.split(r"\n\s*\n", text):
                        cleaned = chunk.strip()
                        if cleaned:
                            section.blocks.append(
                                CanonicalBlock(type="paragraph", text=cleaned)
                            )
            sections.append(section)

        if not sections:
            sections = [CanonicalSection(path=[])]

        return CanonicalDoc(
            doc_title=doc_title or p.stem,
            source_format="pptx",
            source_path=str(p),
            sections=sections,
        )


class PdfParser:
    formats: tuple[str, ...] = ("pdf",)

    def parse(self, path: str | Path, *, doc_title: str | None = None) -> CanonicalDoc:
        import pdfplumber

        p = Path(path)
        sections: list[CanonicalSection] = []
        with pdfplumber.open(str(p)) as pdf:
            for idx, page in enumerate(pdf.pages, start=1):
                section = CanonicalSection(path=[f"Page {idx}"])
                text = page.extract_text() or ""
                for chunk in re.split(r"\n\s*\n", text):
                    cleaned = chunk.strip()
                    if cleaned:
                        section.blocks.append(
                            CanonicalBlock(type="paragraph", text=cleaned, page=idx)
                        )
                for table in page.extract_tables() or []:
                    rows: list[list[str]] = [
                        [("" if cell is None else str(cell)).strip() for cell in row]
                        for row in table
                    ]
                    if rows:
                        tsv = "\n".join("\t".join(r) for r in rows)
                        section.blocks.append(
                            CanonicalBlock(type="table", text=tsv, rows=rows, page=idx)
                        )
                sections.append(section)

        if not sections:
            sections = [CanonicalSection(path=[])]

        return CanonicalDoc(
            doc_title=doc_title or p.stem,
            source_format="pdf",
            source_path=str(p),
            sections=sections,
        )


register(MarkdownParser())
register(TxtParser())
register(DocxParser())
register(PptxParser())
register(PdfParser())

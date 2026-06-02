"""Parser-layer tests: every format -> CanonicalDoc, round-trippable JSON.

Fixtures are built in tmp_path from the real optional libraries (python-docx,
python-pptx, fpdf2). Each is skipped cleanly when its dependency is absent, so the
suite still passes fully offline with whatever subset is installed.
"""

from __future__ import annotations

from pathlib import Path

import pytest

from app.parsing import CanonicalDoc, parse_file


def _assert_canonical(doc: object, *, source_format: str, expected_text: str) -> None:
    """Shared contract assertions for any parsed CanonicalDoc."""
    assert isinstance(doc, CanonicalDoc)
    assert doc.source_format == source_format
    assert len(doc.sections) >= 1
    assert expected_text in doc.to_markdown()
    # JSON round-trip: serialize then validate back to an identical model.
    restored = CanonicalDoc.model_validate_json(doc.model_dump_json())
    assert restored == doc


def test_parse_markdown(tmp_path: Path) -> None:
    md = tmp_path / "policy.md"
    md.write_text(
        "# Policy\n\nTokens expire after 4 hours.\n", encoding="utf-8"
    )
    doc = parse_file(md)
    _assert_canonical(doc, source_format="md", expected_text="Tokens expire after 4 hours.")


def test_parse_html(tmp_path: Path) -> None:
    html = tmp_path / "policy.html"
    html.write_text(
        "<html><body><h1>Policy</h1>"
        "<p>Tokens expire after 4 hours.</p></body></html>",
        encoding="utf-8",
    )
    try:
        doc = parse_file(html)
    except RuntimeError as exc:  # pandoc CLI not installed on this machine
        pytest.skip(f"pandoc unavailable: {exc}")
    _assert_canonical(doc, source_format="html", expected_text="Tokens expire after 4 hours.")


def test_parse_docx(tmp_path: Path) -> None:
    pytest.importorskip("docx")
    import docx

    path = tmp_path / "policy.docx"
    document = docx.Document()
    document.add_heading("Policy", 1)
    document.add_paragraph("Tokens expire after 4 hours.")
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "Key"
    table.cell(0, 1).text = "Value"
    table.cell(1, 0).text = "ttl"
    table.cell(1, 1).text = "4h"
    document.save(str(path))

    doc = parse_file(path)
    _assert_canonical(doc, source_format="docx", expected_text="Tokens expire after 4 hours.")
    # The 2x2 table must survive as structured rows.
    tables = [b for s in doc.sections for b in s.blocks if b.type == "table"]
    assert tables and tables[0].rows is not None and len(tables[0].rows) == 2


def test_parse_pptx(tmp_path: Path) -> None:
    pytest.importorskip("pptx")
    from pptx import Presentation
    from pptx.util import Inches

    path = tmp_path / "deck.pptx"
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[5])  # title-only layout
    slide.shapes.title.text = "Deck"
    box = slide.shapes.add_textbox(Inches(1), Inches(2), Inches(6), Inches(2))
    box.text_frame.text = "Tokens expire after 4 hours."
    prs.save(str(path))

    doc = parse_file(path)
    _assert_canonical(doc, source_format="pptx", expected_text="Tokens expire after 4 hours.")


def test_parse_pdf(tmp_path: Path) -> None:
    pytest.importorskip("fpdf")
    # The native PDF parser reads via pdfplumber; both libs are required end-to-end.
    pytest.importorskip("pdfplumber")
    from fpdf import FPDF

    path = tmp_path / "audit.pdf"
    pdf = FPDF()
    pdf.add_page()
    pdf.set_font("helvetica", size=12)
    pdf.multi_cell(0, 10, "Audit logs retained for 2 years.")
    pdf.output(str(path))

    doc = parse_file(path)
    _assert_canonical(doc, source_format="pdf", expected_text="Audit logs retained for 2 years.")
